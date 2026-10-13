#!/usr/bin/env python3
"""Liest eine (mit LibreOffice neu berechnete) Mappe, sucht Fehlerwerte und gibt Kernergebnisse aus."""
import sys
from openpyxl import load_workbook

ERR = ("#N/A", "#NV", "#VALUE!", "#WERT!", "#DIV/0!", "#REF!", "#BEZUG!", "#NAME?", "#NUM!", "#ZAHL!", "#NULL!", "Err:", "#SPILL")


def load(path):
    wb = load_workbook(path, data_only=True)
    errs = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith(ERR):
                    errs.append(f"{ws.title}!{c.coordinate}={c.value}")
    names = {}
    for n, dn in wb.defined_names.items():
        try:
            dests = list(dn.destinations)
        except Exception:
            continue
        if len(dests) != 1:
            continue
        sh, ref = dests[0]
        ref = ref.replace("$", "")
        if ":" in ref:
            names[n] = [c.value for row in wb[sh][ref] for c in row]
        else:
            names[n] = wb[sh][ref].value
    return wb, names, errs


if __name__ == "__main__":
    wb, N, errs = load(sys.argv[1])
    print("Fehlerzellen:", len(errs), errs[:20])
    keys = sys.argv[2:] or ["Modus", "Szenario", "Kaufpreis", "NK_Summe", "NK_Quote", "Gesamtkosten", "Gebaeudeanteil", "AK_Gebaeude",
                            "Kredit1_Betrag", "Darlehen_Summe", "EK_Einsatz", "Grenzsteuersatz_calc", "AfA_Art", "AfA_BMG_Gebaeude",
                            "Brutto_Rendite", "Netto_Rendite", "KP_Faktor", "CF_vor_1_mtl", "Steuer_1", "CF_nach_1_mtl", "EK_Rendite_1",
                            "IRR_Wert", "DSCR_1", "BE_V_Jahr", "BE_CF_Jahr", "Vermoegen_N", "EndA", "EndB", "Vorteil", "Score", "Fazit"]
    for k in keys:
        print(f"{k:24s} {N.get(k)}")

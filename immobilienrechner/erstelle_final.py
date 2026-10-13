#!/usr/bin/env python3
"""Baut die Mappe, rechnet eine Kopie mit LibreOffice, prüft sie und schreibt die Ergebnisse als
zwischengespeicherte Werte in die (unveränderte) openpyxl-Datei. Excel rechnet beim Öffnen ohnehin neu
(fullCalcOnLoad); die Werte sorgen dafür, dass Vorschauen/Geschützte Ansicht sofort Zahlen zeigen."""
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from xml.sax.saxutils import escape

from openpyxl import load_workbook

import build_immobilienrechner as B
from pruefe_rechner import load

RECALC = os.environ.get("RECALC", "/root/.claude/skills/synced/80589829-e181-4b8f-8d3b-74fa4a5ef02c_9c8e0ddd-49e5-487f-8836-8bf3f0d3ab68/xlsx/scripts/recalc.py")


def sheet_files(z):
    wbx = z.read("xl/workbook.xml").decode()
    rels = z.read("xl/_rels/workbook.xml.rels").decode()
    rmap = dict(re.findall(r'<Relationship[^>]*Id="(rId\d+)"[^>]*Target="([^"]+)"', rels))
    rmap.update({a: b for b, a in re.findall(r'<Relationship[^>]*Target="([^"]+)"[^>]*Id="(rId\d+)"', rels)})
    out = {}
    for name, rid in re.findall(r'<sheet[^>]*name="([^"]+)"[^>]*r:id="(rId\d+)"', wbx):
        t = rmap[rid].lstrip("/")
        out[name.replace("&amp;", "&")] = t if t.startswith("xl/") else "xl/" + t
    return out


def inject(src, calc, dst):
    vals = load_workbook(calc, data_only=True)
    zin = zipfile.ZipFile(src)
    files = sheet_files(zin)
    target = {v: k for k, v in files.items()}
    pat = re.compile(r'<c r="([A-Z]+\d+)"((?: s="\d+")?)><f>(.*?)</f><v\s*/?>(?:</v>)?</c>')
    n = 0
    with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename in target:
                ws = vals[target[item.filename]]
                xml = data.decode()

                def rep(mm):
                    nonlocal n
                    ref, s, f = mm.groups()
                    v = ws[ref].value
                    n += 1
                    if isinstance(v, bool):
                        return f'<c r="{ref}"{s} t="b"><f>{f}</f><v>{int(v)}</v></c>'
                    if isinstance(v, (int, float)):
                        return f'<c r="{ref}"{s}><f>{f}</f><v>{repr(float(v)) if isinstance(v, float) else v}</v></c>'
                    if v is None:
                        v = ""
                    return f'<c r="{ref}"{s} t="str"><f>{f}</f><v>{escape(str(v))}</v></c>'

                xml = pat.sub(rep, xml)
                data = xml.encode()
            zout.writestr(item, data)
    return n


def main(out):
    tmp = tempfile.mkdtemp()
    raw = os.path.join(tmp, "roh.xlsx")
    calc = os.path.join(tmp, "berechnet.xlsx")
    B.build(raw)
    shutil.copy(raw, calc)
    res = subprocess.run([sys.executable, RECALC, calc, "300"], capture_output=True, text=True)
    print(res.stdout.strip())
    assert '"status": "success"' in res.stdout, "Neuberechnung fehlgeschlagen"
    _, _, errs = load(calc)
    assert not errs, errs[:10]
    n = inject(raw, calc, out)
    # Kontrolle: Datei lesbar, Werte vorhanden, Formeln erhalten
    wb_f = load_workbook(out)
    wb_v = load_workbook(out, data_only=True)
    assert wb_f["Kennzahlen & Ampel"].protection.sheet
    f = sum(1 for ws in wb_f.worksheets for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("="))
    print(f"Werte eingesetzt: {n}, Formeln in Ausgabedatei: {f}")
    _, N, errs = load(out)
    assert not errs
    print("Fazit:", N["Fazit"], "| IRR:", N["IRR_Wert"])
    shutil.rmtree(tmp)


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "Immobilienrechner_2026.xlsx")
